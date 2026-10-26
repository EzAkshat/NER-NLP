import io
import zipfile

import pytest
from docx import Document
from helpers import make_pdf

from ner_app.errors import (
    EmptyDocumentError,
    ExtractionError,
    ResourceLimitError,
)
from ner_app.ingestion import Limits, extract_document


def codes(doc):
    return {w.code for w in doc.warnings}


def save_docx(tmp_path, name, build):
    document = Document()
    build(document)
    path = tmp_path / name
    document.save(path)
    return path


# ---------- DOCX ----------

def test_docx_paragraphs_and_tables_in_order(tmp_path):
    def build(d):
        d.add_paragraph("Intro paragraph.")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Org"
        table.cell(1, 0).text = "Jane Smith"
        table.cell(1, 1).text = "Example Technologies"
        d.add_paragraph("Closing paragraph.")

    doc = extract_document(save_docx(tmp_path, "a.docx", build))
    text = doc.text
    assert doc.detected_format == "docx"
    assert "Name\nOrg\nJane Smith\nExample Technologies" in text
    assert text.index("Intro") < text.index("Example Technologies") < text.index("Closing")


def test_docx_unicode(tmp_path):
    path = save_docx(tmp_path, "u.docx", lambda d: d.add_paragraph("Caf\u00e9 M\u00fcller hosted guests."))
    assert "Caf\u00e9 M\u00fcller hosted guests." in extract_document(path).text


def test_docx_without_text_is_empty(tmp_path):
    with pytest.raises(EmptyDocumentError):
        extract_document(save_docx(tmp_path, "e.docx", lambda d: None))


def test_docx_not_a_zip(tmp_path):
    path = tmp_path / "fake.docx"
    path.write_bytes(b"this is not a zip file")
    with pytest.raises(ExtractionError):
        extract_document(path)


def test_docx_zip_but_not_word(tmp_path):
    path = tmp_path / "other.docx"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    path.write_bytes(buf.getvalue())
    with pytest.raises(ExtractionError):
        extract_document(path)


def test_docx_macros_are_ignored_with_warning(tmp_path):
    path = save_docx(tmp_path, "m.docx", lambda d: d.add_paragraph("Hello world."))
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("word/vbaProject.bin", b"not a real macro")
    doc = extract_document(path)
    assert "Hello world." in doc.text
    assert "docx_macros_ignored" in codes(doc)


def test_docx_compression_bomb_rejected(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * 5_000_000)
    path = tmp_path / "bomb.docx"
    path.write_bytes(buf.getvalue())
    with pytest.raises(ResourceLimitError):
        extract_document(path)


def test_docx_uncompressed_size_limit(tmp_path):
    path = save_docx(tmp_path, "s.docx", lambda d: d.add_paragraph("Hello world."))
    with pytest.raises(ResourceLimitError):
        extract_document(path, Limits(max_uncompressed_bytes=1000))


def test_docx_entry_count_limit(tmp_path):
    path = save_docx(tmp_path, "n.docx", lambda d: d.add_paragraph("Hello world."))
    with pytest.raises(ResourceLimitError):
        extract_document(path, Limits(max_archive_entries=2))


# ---------- PDF ----------

def test_pdf_text_extraction_and_page_order(tmp_path):
    path = tmp_path / "a.pdf"
    path.write_bytes(make_pdf(["Jane Smith visited Berlin.", "Second page about Microsoft."]))
    doc = extract_document(path)
    assert doc.detected_format == "pdf"
    assert doc.text.index("Jane Smith visited Berlin.") < doc.text.index("Second page about Microsoft.")


def test_pdf_blank_page_produces_warning(tmp_path):
    path = tmp_path / "b.pdf"
    path.write_bytes(make_pdf(["Jane Smith visited Berlin.", ""]))
    doc = extract_document(path)
    assert "Jane Smith visited Berlin." in doc.text
    assert "pdf_pages_without_text" in codes(doc)


def test_pdf_with_no_text_is_reported_as_scanned(tmp_path):
    path = tmp_path / "scan.pdf"
    path.write_bytes(make_pdf(["", ""]))
    with pytest.raises(EmptyDocumentError, match="scanned"):
        extract_document(path)


def test_pdf_garbage_rejected(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"%PDF-1.4\nthis is not really a pdf")
    with pytest.raises(ExtractionError):
        extract_document(path)


def test_pdf_page_limit(tmp_path):
    path = tmp_path / "p.pdf"
    path.write_bytes(make_pdf(["One.", "Two."]))
    with pytest.raises(ResourceLimitError):
        extract_document(path, Limits(max_pdf_pages=1))


def test_pdf_char_limit(tmp_path):
    path = tmp_path / "c.pdf"
    path.write_bytes(make_pdf(["Jane Smith visited Berlin."]))
    with pytest.raises(ResourceLimitError):
        extract_document(path, Limits(max_chars=5))